"""Build only the dedicated VA resume from its shared factual source."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
data = json.loads((ROOT / "scripts/va-resume.json").read_text())
output = ROOT / "downloads/Eric_John_Cruz_Virtual_Assistant_Resume.docx"
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.52)
sec.left_margin = sec.right_margin = Inches(.65)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.02
# Remove inherited title rules and character spacing from the default template.
for style in doc.styles:
    for tag in ["w:pBdr", "w:spacing"]:
        if tag == "w:pBdr":
            targets = style.element.xpath(".//w:pPr/w:pBdr")
        else:
            targets = style.element.xpath(".//w:rPr/w:spacing")
        for node in targets:
            node.getparent().remove(node)
doc.styles["Subtitle"].font.italic = False
doc.styles["Normal"].font.size = Pt(10)
doc.styles["Title"].font.size = Pt(23)
doc.styles["Title"].font.bold = True
doc.styles["Subtitle"].font.size = Pt(11)
doc.styles["Heading 1"].font.size = Pt(11)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
doc.styles["Heading 2"].font.size = Pt(10)
doc.styles["Heading 2"].font.bold = True
doc.styles["Heading 2"].paragraph_format.space_before = Pt(4)
doc.styles["List Bullet"].font.size = Pt(10)
doc.styles["List Bullet"].paragraph_format.space_after = Pt(3)
doc.styles["List Bullet"].paragraph_format.left_indent = Inches(.15)
doc.styles["List Bullet"].paragraph_format.first_line_indent = Inches(-.12)

def link(p, label, url):
    item = OxmlElement("w:hyperlink")
    item.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "000000"); props.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); props.append(underline)
    run.append(props)
    text = OxmlElement("w:t"); text.text = label; run.append(text)
    item.append(run); p._p.append(item)

def para(text="", bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix + ": ").bold = True
    p.add_run(text)
    return p

doc.add_paragraph(data["name"], "Title")
doc.add_paragraph(data["headline"] + " | " + data["subtitle"], "Subtitle")
p = para(data["location"] + " | ")
link(p, data["email"], "mailto:" + data["email"])
p = para()
link(p, "Portfolio: ericcruzdigital.github.io/virtual-assistant/", data["portfolio"])
p.add_run("  |  ")
link(p, "LinkedIn", data["linkedin"])
doc.add_heading("Professional Summary", 1)
para(data["summary"])
doc.add_heading("Core Skills", 1)
for label, text in data["skills"]:
    para(text, label)
doc.add_heading("Professional Experience", 1)
for item in data["experience"]:
    doc.add_heading(item["title"], 2)
    p = para(item["company"] + " | " + item["dates"])
    p.paragraph_format.keep_with_next = True
    for bullet in item["bullets"]:
        doc.add_paragraph(bullet, "List Bullet")
    if item.get("engagements"):
        para(item["engagements"])
doc.add_heading("Tools", 1)
for label, text in data["tools"]:
    para(text + ".", label)
doc.add_heading("Selected Work", 1)
para(data["work"])
p = para()
link(p, data["portfolio"] + "work.html", data["portfolio"] + "work.html")
doc.core_properties.author = data["name"]
doc.core_properties.title = "Eric John Cruz Virtual Assistant Resume"
doc.core_properties.subject = "Marketing operations and administrative support"
doc.core_properties.keywords = "Virtual Assistant, Marketing, Operations, Reporting"
doc.core_properties.comments = ""
doc.core_properties.last_modified_by = data["name"]
output.parent.mkdir(exist_ok=True)
doc.save(output)
print(output.relative_to(ROOT))
